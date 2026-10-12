#!/usr/bin/env python3
"""deck.py - build the fully editable scene-by-scene deck.

Not pictures of slides, and not a flat plate with text on top: every element on every frame is its
own native PowerPoint object.

  background   one rectangle with the real 3-stop linear gradient
  cards, pills, bars, rules, dots   rounded rectangles / freeforms with solid fill, real stroke
                                    (dashes stay dashes), real soft shadow
  Figma artwork  the real SVG embedded as an SVG picture with a PNG fallback, so PowerPoint's
                 "Graphics Format > Convert to Shape" turns each one into editable vectors
  every line of copy  a real text box at the measured position, size, colour and weight

Run `node shapes.mjs` first; it writes _shapes/<name>.json and _shapes/<name>-NN.svg.
Set BRAND_FONT below; that font must be installed locally for the deck to render on brand.
"""
import os, re, json, glob, subprocess, hashlib

BRAND_FONT = os.environ.get("BRAND_FONT", "Inter Tight")   # one variable, swap once
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree
from scenes import SCENES, SLUG, TITLE

B = os.path.dirname(os.path.abspath(__file__))
JOB = os.path.dirname(os.path.dirname(B))   # build/deck/ -> the job root
SH = os.path.join(B, "_shapes")
PNG = os.path.join(B, "_shapes", "_png")
os.makedirs(PNG, exist_ok=True)
OUT = os.path.join(JOB, "out", f"{SLUG}_EDITABLE.pptx")

STAGE_W, STAGE_H = 1920, 1080
EMU_PX = 6350                      # 1 stage px = 1/144 in = 6350 EMU
PT_PX = 0.5                        # 1 stage px = 0.5 pt at this slide size
E = lambda v: Emu(int(round(v * EMU_PX)))
ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT,
         "justify": PP_ALIGN.JUSTIFY, "start": PP_ALIGN.LEFT, "end": PP_ALIGN.RIGHT}
DASH = {"dashed": MSO_LINE_DASH_STYLE.DASH, "dotted": MSO_LINE_DASH_STYLE.ROUND_DOT}
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
RNS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
SVGNS = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
# the .stage background, straight out of gen.py's HEAD
BG_STOPS = [(0, "EEF2FF"), (52, "F8FAFC"), (100, "FFF7ED")]   # per job: the frame background stops

CAP = {n: (label, beat) for n, label, beat in SCENES}
_svg_seq = [0]


def bg_at(y):
    """The background gradient colour at a given stage y, used to flatten translucent fills."""
    t = max(0.0, min(1.0, y / STAGE_H)) * 100
    for (p0, c0), (p1, c1) in zip(BG_STOPS, BG_STOPS[1:]):
        if t <= p1:
            f = 0 if p1 == p0 else (t - p0) / (p1 - p0)
            a = [int(c0[i:i+2], 16) for i in (0, 2, 4)]
            b = [int(c1[i:i+2], 16) for i in (0, 2, 4)]
            return [a[i] + (b[i] - a[i]) * f for i in range(3)]
    c = BG_STOPS[-1][1]
    return [int(c[i:i+2], 16) for i in (0, 2, 4)]


def flatten(hexv, alpha, y):
    """LibreOffice, Google Slides and Keynote all disagree about a:alpha on a solid fill (verified:
    LO renders 42% white as a grey). Every translucent fill on this board is a flat wash over the
    background gradient, so composite it once here and ship a solid colour that every renderer
    agrees on."""
    fg = [int(hexv[i:i+2], 16) for i in (0, 2, 4)]
    bg = bg_at(y)
    return "".join(f"{int(round(fg[i] * alpha + bg[i] * (1 - alpha))):02X}" for i in range(3))


# --- low level OOXML helpers -------------------------------------------------------------
def _sp_pr(shape):
    return shape._element.spPr


def strip_theme(shape):
    """Drop the <p:style> theme reference python-pptx attaches to every autoshape. An empty
    <a:effectLst/> is supposed to override <a:effectRef>, and PowerPoint honours that, but
    LibreOffice does not: it draws the theme drop shadow anyway. With no <p:style> at all every
    renderer has to use spPr, and the shapes stay clean to restyle by hand."""
    el = shape._element
    st = el.find(qn("p:style"))
    if st is not None:
        el.remove(st)


def set_alpha(shape, alpha):
    """Fill transparency. python-pptx writes a bare srgbClr; add the alpha child ourselves."""
    if alpha >= 0.999:
        return
    clr = _sp_pr(shape).find(qn("a:solidFill"))
    if clr is None:
        return
    srgb = clr.find(qn("a:srgbClr"))
    if srgb is None:
        return
    a = etree.SubElement(srgb, qn("a:alpha"))
    a.set("val", str(int(round(alpha * 100000))))


def add_shadow(shape, css, opacity=1.0):
    """Translate a CSS box-shadow into a real PowerPoint outer shadow."""
    col = re.search(r"rgba?\(([\d.]+),\s*([\d.]+),\s*([\d.]+)(?:,\s*([\d.]+))?\)", css)
    nums = re.findall(r"(-?[\d.]+)px", css)
    if not col or len(nums) < 3:
        return
    r, g, b = (int(float(col.group(i))) for i in (1, 2, 3))
    alpha = float(col.group(4) or 1)
    dx, dy, blur = (float(n) for n in nums[:3])
    dist = (dx ** 2 + dy ** 2) ** 0.5
    spPr = _sp_pr(shape)
    lst = etree.SubElement(spPr, qn("a:effectLst"))
    sh = etree.SubElement(lst, qn("a:outerShdw"))
    sh.set("blurRad", str(int(blur * EMU_PX)))
    sh.set("dist", str(int(dist * EMU_PX)))
    sh.set("dir", "5400000")                      # straight down, as every shadow here is
    sh.set("rotWithShape", "0")
    c = etree.SubElement(sh, qn("a:srgbClr"))
    c.set("val", f"{r:02X}{g:02X}{b:02X}")
    al = etree.SubElement(c, qn("a:alpha"))
    al.set("val", str(int(max(0.0, min(1.0, alpha * opacity)) * 100000)))


def gradient_bg(slide):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, E(STAGE_W), E(STAGE_H))
    sp.name = "background"
    sp.line.fill.background()
    sp.shadow.inherit = False
    strip_theme(sp)
    spPr = _sp_pr(sp)
    for tag in ("a:solidFill", "a:gradFill", "a:noFill"):
        old = spPr.find(qn(tag))
        if old is not None:
            spPr.remove(old)
    grad = etree.Element(qn("a:gradFill"))
    grad.set("rotWithShape", "1")
    lst = etree.SubElement(grad, qn("a:gsLst"))
    for pos, hexv in BG_STOPS:
        gs = etree.SubElement(lst, qn("a:gs"))
        gs.set("pos", str(int(pos * 1000)))
        etree.SubElement(gs, qn("a:srgbClr")).set("val", hexv)
    lin = etree.SubElement(grad, qn("a:lin"))
    lin.set("ang", "2700000")                     # 135deg CSS = top-left to bottom-right (gen.py stage)
    lin.set("scaled", "0")
    ln = spPr.find(qn("a:ln"))
    spPr.insert(list(spPr).index(ln) if ln is not None else len(spPr), grad)
    return sp


def rounded_points(x, y, w, h, radii, seg=8):
    """A rounded rectangle with per-corner radii, as a polyline (for the corner bars and the
    clipped header rows, where the four radii differ and roundRect cannot express it)."""
    import math
    r = [max(0.0, min(v, min(w, h) / 2)) for v in radii]          # tl, tr, br, bl
    pts, corners = [], [
        (x + w - r[1], y + r[1], r[1], -90, 0),
        (x + w - r[2], y + h - r[2], r[2], 0, 90),
        (x + r[3], y + h - r[3], r[3], 90, 180),
        (x + r[0], y + r[0], r[0], 180, 270),
    ]
    for cx, cy, rad, a0, a1 in corners:
        if rad <= 0.01:
            pts.append((cx, cy)); continue
        for i in range(seg + 1):
            a = math.radians(a0 + (a1 - a0) * i / seg)
            pts.append((cx + rad * math.cos(a), cy + rad * math.sin(a)))
    return pts


def add_box(slide, it):
    x, y, w, h = it["x"], it["y"], it["w"], it["h"]
    radii = it.get("radii") or [0, 0, 0, 0]
    uniform = max(radii) - min(radii) < 0.6
    # A CSS border-radius:50% box is an ELLIPSE. Emitting it as a roundRect with adjustment 0.5
    # draws a squircle in LibreOffice (and reads wrong in PowerPoint too) - caught on the
    # "data leaves your control" frame, where all three icon discs came back as rounded squares.
    # ...but only when the box is SQUARE. On a 27x14 pill, radius 7 is a stadium with straight
    # sides, not an ellipse - the un-squared test turned a whole grid of dashes into ovals.
    circular = uniform and min(radii) >= min(w, h) / 2 - 0.6 and abs(w - h) <= 1.5
    if circular:
        sp = slide.shapes.add_shape(MSO_SHAPE.OVAL, E(x), E(y), E(w), E(h))
    elif max(radii) < 0.6:
        sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, E(x), E(y), E(w), E(h))
    elif uniform:
        sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, E(x), E(y), E(w), E(h))
        sp.adjustments[0] = min(0.5, radii[0] / max(1.0, min(w, h)))
    else:
        pts = rounded_points(x, y, w, h, radii)
        bld = slide.shapes.build_freeform(E(pts[0][0]), E(pts[0][1]))
        bld.add_line_segments([(E(px), E(py)) for px, py in pts[1:]], close=True)
        sp = bld.convert_to_shape()
    sp.shadow.inherit = False
    strip_theme(sp)
    if it.get("fill"):
        a = it["fill"]["a"] * it.get("opacity", 1)
        sp.fill.solid()
        sp.fill.fore_color.rgb = RGBColor.from_string(
            it["fill"]["hex"] if a >= 0.999 else flatten(it["fill"]["hex"], a, y + h / 2))
    else:
        sp.fill.background()
    st = it.get("stroke")
    op = it.get("opacity", 1)
    if st and st["w"] > 0:
        # A faded element fades its border and its shadow too, not just its fill. Keyed on the fill
        # alone, a ghosted card keeps a full-strength outline and a full-strength shadow over a
        # nearly-background fill, so it reads as an empty box instead of fading out.
        sp.line.color.rgb = RGBColor.from_string(
            st["hex"] if op >= 0.999 else flatten(st["hex"], op, y + h / 2))
        sp.line.width = E(st["w"])
        if st["dash"] in DASH:
            sp.line.dash_style = DASH[st["dash"]]
    else:
        sp.line.fill.background()
    if it.get("shadow"):
        add_shadow(sp, it["shadow"], op)
    return sp


def add_img(slide, path, it):
    """Place a real bitmap brand asset (some brand icons ship only as PNG). Straight picture: there
    is no vector to convert to, so it stays an image - flagged in the job NOTES."""
    return slide.shapes.add_picture(path, E(it["x"]), E(it["y"]), E(it["w"]), E(it["h"]))


def add_svg(slide, path, it):
    """Place the real Figma vector: an SVG picture with a PNG fallback, so PowerPoint can convert
    it to editable shapes and every other renderer still shows the artwork."""
    w, h = it["w"], it["h"]
    key = hashlib.md5((os.path.basename(path) + f"{w:.1f}x{h:.1f}").encode()).hexdigest()[:12]
    png = os.path.join(PNG, key + ".png")
    if not os.path.exists(png):
        subprocess.run(["rsvg-convert", "-w", str(max(8, int(w * 4))), "-h", str(max(8, int(h * 4))),
                        "-f", "png", "-o", png, path], check=True)
    pic = slide.shapes.add_picture(png, E(it["x"]), E(it["y"]), E(w), E(h))
    _svg_seq[0] += 1
    part = Part(PackURI(f"/ppt/media/vector{_svg_seq[0]}.svg"), "image/svg+xml",
                slide.part.package, open(path, "rb").read())
    rId = slide.part.relate_to(part, RT.IMAGE)
    blip = pic._element.blipFill.find(qn("a:blip"))
    ext_lst = etree.SubElement(blip, qn("a:extLst"))
    ext = etree.SubElement(ext_lst, qn("a:ext"))
    ext.set("uri", "{96DAC541-7B7A-43D3-8B79-37D633B846F1}")
    sb = etree.SubElement(ext, f"{{{SVGNS}}}svgBlip")
    sb.set(f"{{{RNS}}}embed", rId)
    op = it.get("opacity", 1)
    if op < 0.999:
        amf = etree.SubElement(blip, qn("a:alphaModFix"))
        amf.set("amt", str(int(op * 100000)))
        blip.remove(ext_lst); blip.append(ext_lst)      # extLst must stay last
    return pic


def add_text(slide, b):
    x, y, w = b["x"], b["y"], b["w"]
    align = b["align"] if b["align"] in ALIGN else "left"
    slack = 300
    if align == "center":
        w += slack * 2; x -= slack
    elif align == "right":
        w += slack; x -= slack
    else:
        w = min(w + slack, STAGE_W - x)
    w = max(w, 40)
    size, pitch = b["size"], b["pitch"]
    # Anchor the block in the MIDDLE of a box that is exactly the measured text block. Every engine
    # (PowerPoint, LibreOffice, Slides, Keynote) hangs the first line off the top of a box using its
    # own ascent, and they disagree; centring inside the real block cancels that difference instead
    # of guessing at one engine's half-leading.
    pad = 120.0
    tb = slide.shapes.add_textbox(E(x), E(y - pad), E(w), E(b["h"] + pad * 2))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, runs in enumerate(b["lines"]):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = ALIGN[align]
        if len(b["lines"]) > 1 and size:
            ratio = round(pitch / size, 3)
            if 0.1 <= ratio <= 132.0:
                para.line_spacing = ratio
        for r in runs:
            run = para.add_run()
            t = r["t"]
            run.text = ((" " if t[:1].isspace() else "") + " ".join(t.split())
                        + (" " if t[-1:].isspace() and t.strip() else ""))
            f = run.font
            f.name = "Menlo" if r.get("mono", b["mono"]) else BRAND_FONT   # per run: an
            # inline <span class="mono"> inside a brand-font line keeps its mono face
            f.size = Pt(round(size * PT_PX, 1))
            f.bold = bool(r["bold"])
            f.color.rgb = RGBColor.from_string(r["color"].upper())
    return tb


# --- build -------------------------------------------------------------------------------
prs = Presentation()
prs.slide_width = Emu(STAGE_W * EMU_PX)
prs.slide_height = Emu(STAGE_H * EMU_PX)
blank = prs.slide_layouts[6]
counts = []

for name, label, beat in SCENES:
    items = json.load(open(os.path.join(SH, name + ".json")))
    slide = prs.slides.add_slide(blank)
    gradient_bg(slide)
    n = {"box": 0, "svg": 0, "img": 0, "text": 0}
    for it in items:
        if it["kind"] == "box":
            add_box(slide, it)
        elif it["kind"] == "svg":
            add_svg(slide, os.path.join(SH, f"{name}-{it['idx']:02d}.svg"), it)
        elif it["kind"] == "img":
            cands = glob.glob(os.path.join(SH, f"{name}-img-{it['idx']:02d}.*"))
            if not cands:
                raise SystemExit(f"missing bitmap for img {it['idx']} of {name}")
            add_img(slide, cands[0], it)
        else:
            add_text(slide, it)
        n[it["kind"]] += 1
    counts.append((name, n))
    notes = slide.notes_slide.notes_text_frame
    notes.text = (f"{name}  |  {label}\n\n{beat}\n\n"
                  f"Fully editable: 1 gradient background, {n['box']} shapes, {n['svg']} Figma vectors, "
                  f"{n['img']} bitmap brand icons, {n['text']} text boxes. The vectors are SVG pictures: select one and use "
                  f"Graphics Format > Convert to Shape to edit its paths. Type is {BRAND_FONT}.")

os.makedirs(os.path.dirname(OUT), exist_ok=True)
prs.save(OUT)
tot = {k: sum(c[k] for _, c in counts) for k in ("box", "svg", "text")}
print(f"wrote {OUT}")
print(f"  {len(counts)} slides, {tot['box']} shapes, {tot['svg']} vectors, {tot['text']} text boxes, "
      f"{os.path.getsize(OUT)/1e6:.1f} MB")
